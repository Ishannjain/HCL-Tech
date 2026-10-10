import re
from pathlib import Path

def parse_file(filename, content):
    ext = Path(filename).suffix.lower()
    if ext in {'.txt','.md'}:
        return [(1, content.decode('utf-8-sig'))]
    if ext == '.pdf':
        from io import BytesIO
        from pypdf import PdfReader
        pages = [(i+1,p.extract_text() or '') for i,p in enumerate(PdfReader(BytesIO(content)).pages)]
        if not any(text.strip() for _,text in pages):
            raise ValueError('Scanned PDF needs OCR. Supply searchable PDF or extracted text; this file was not activated.')
        return pages
    if ext == '.docx':
        from docx import Document
        from io import BytesIO
        doc = Document(BytesIO(content))
        parts = [p.text for p in doc.paragraphs]
        parts += [' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows]
        return [(1,'\n'.join(parts))]
    raise ValueError('Supported document formats: TXT, MD, searchable PDF, DOCX')

def make_chunks(pages, max_chars=1500):
    result = []
    for page,text in pages:
        section, block = f'Page {page}', []
        def emit():
            body = '\n'.join(block).strip()
            if body:
                remaining=body
                while remaining:
                    end=min(len(remaining),max_chars)
                    if end<len(remaining):
                        boundary=remaining.rfind(' ',0,end)
                        if boundary>0:end=boundary
                    result.append({'section':section,'page':page,'text':remaining[:end]})
                    remaining=remaining[end:].lstrip()
        for line in text.splitlines():
            match = re.match(r'^#{1,6}\s+(.+?)\s*$',line)
            if match:
                emit()
                section,block = match.group(1),[]
            else:
                block.append(line)
        emit()
    return result
